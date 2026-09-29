"""Сводная по заявке и счетам поставщиков без нейросетевых расчётов."""

from __future__ import annotations

import hashlib
import json
import os
import re
import subprocess
import tempfile
from collections import defaultdict
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import Iterable


MONEY = Decimal("0.01")


def decimal_value(value: object) -> Decimal:
    if isinstance(value, Decimal):
        return value
    if isinstance(value, (int, float)):
        return Decimal(str(value))
    raw = str(value or "").strip()
    raw = raw.replace("\u00a0", "").replace(" ", "")
    raw = raw.replace("о", "0").replace("О", "0").replace("o", "0").replace("O", "0")
    if "," in raw and "." in raw:
        raw = raw.replace(",", "") if raw.rfind(".") > raw.rfind(",") else raw.replace(".", "").replace(",", ".")
    elif "," in raw:
        raw = raw.replace(",", ".")
    raw = re.sub(r"[^0-9.\-]", "", raw)
    try:
        return Decimal(raw)
    except (InvalidOperation, ValueError):
        raise ValueError(f"Не удалось прочитать число: {value!r}") from None


def money(value: Decimal) -> Decimal:
    return value.quantize(MONEY, rounding=ROUND_HALF_UP)


def normalized_name(value: str) -> str:
    name = value.upper().replace("Х", "X").replace("×", "X")
    name = re.sub(r"(?<=\d)\s*['’`´]\s*(?=\d)", "X", name)
    return re.sub(r"\s+", " ", name).strip()


def product_key(value: str) -> tuple[str, tuple[str, ...]] | None:
    name = normalized_name(value)
    if "ПРОФ" in name and "ТРУБ" in name:
        match = re.search(r"(\d{2,3})\s*X\s*(\d{2,3})\s*X\s*(\d+(?:[.,]\d+)?)", name)
        return ("профильная труба", tuple(_dimension(part) for part in match.groups())) if match else None
    if "ШВЕЛЛЕР" in name:
        match = re.search(r"ШВЕЛЛЕР[^\d]*(\d{1,2})\s*([ПУ])", name)
        return ("швеллер", (match[1], match[2])) if match else None
    if "БАЛК" in name or "ДВУТАВР" in name:
        match = re.search(r"(\d{2})\s*([БК])\s*(\d)", name)
        return ("балка", match.groups()) if match else None
    if "УГОЛ" in name:
        match = re.search(r"(\d{2,3})\s*X\s*(\d{1,3})(?:\s*X\s*(\d+(?:[.,]\d+)?))?", name)
        if match:
            parts = match.groups()
            return ("уголок", tuple(_dimension(part) for part in
                    (parts if parts[2] else (parts[0], parts[0], parts[1]))))
    if "КРУГЛЯК" in name or re.search(r"\bКРУГ\b", name):
        match = re.search(r"(?:КРУГЛЯК|КРУГ)\s*(?:Ф|Ø)?\s*(\d+(?:[.,]\d+)?)", name)
        return ("круг", (_dimension(match[1]),)) if match else None
    if "ТРУБ" in name and "ПРОФ" not in name:
        match = re.search(r"(?:Ф|Ø)\s*(\d+(?:[.,]\d+)?)|КРУГЛАЯ[^\d]*(\d+(?:[.,]\d+)?)", name)
        return ("круглая труба", (_dimension(next(part for part in match.groups() if part)),)) if match else None
    return None


def _dimension(value: str) -> str:
    number = decimal_value(value)
    return format(number.normalize(), "f")


def grade_tokens(value: str) -> tuple[str, ...]:
    name = normalized_name(value).replace("C", "С")
    return tuple(sorted(set(re.findall(r"С\s*([2345]\d{2})", name))))


def length_from_name(value: str) -> Decimal | None:
    # В счетах встречаются (12000), (12), (12 м), (11,7).
    for match in re.finditer(r"\(\s*(\d{1,5}(?:[.,]\d+)?)", normalized_name(value)):
        number = decimal_value(match[1])
        if number >= 1000:
            return number / 1000
        if Decimal("0.1") <= number <= 30:
            return number
    return None


@dataclass(frozen=True)
class RequestLine:
    source_row: int
    item_no: int
    name: str
    length_m: Decimal
    quantity_pcs: int
    grade: str

    @property
    def key(self) -> tuple[str, tuple[str, ...]] | None:
        return product_key(self.name)


@dataclass(frozen=True)
class OfferSource:
    path: Path
    supplier: str


@dataclass(frozen=True)
class InvoiceLine:
    supplier: str
    invoice_number: str
    source_file: Path
    page: int
    row_number: int
    name: str
    quantity: Decimal
    unit: str
    mass_t: Decimal | None
    amount: Decimal
    price_per_t: Decimal | None
    length_m: Decimal | None
    grades: tuple[str, ...]
    warehouse: str
    notes: tuple[str, ...] = ()

    @property
    def key(self) -> tuple[str, tuple[str, ...]] | None:
        return product_key(self.name)


@dataclass(frozen=True)
class Invoice:
    supplier: str
    number: str
    source_file: Path
    lines: tuple[InvoiceLine, ...]
    issues: tuple[str, ...] = ()


@dataclass(frozen=True)
class Allocation:
    line: InvoiceLine
    request: RequestLine | None
    status: str


@dataclass(frozen=True)
class CombinedOffer:
    supplier: str
    request: RequestLine
    source_lines: tuple[InvoiceLine, ...]
    mass_t: Decimal | None
    amount: Decimal
    price_per_t: Decimal | None
    status: str


@dataclass
class SummaryResult:
    request_lines: list[RequestLine]
    suppliers: list[str]
    allocations: list[Allocation]
    offers: dict[tuple[int, str], CombinedOffer]
    issues: list[str]
    output_path: Path

    def find_offer(self, *, item_no: int, length_m: Decimal, supplier: str) -> CombinedOffer | None:
        for request in self.request_lines:
            if request.item_no == item_no and request.length_m == length_m:
                return self.offers.get((request.source_row, supplier))
        return None


def read_request(path: Path) -> list[RequestLine]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("Для сводной установите зависимости: pip install -r requirements.txt") from exc
    workbook = load_workbook(path, read_only=False, data_only=True)
    sheet = workbook.active
    lines: list[RequestLine] = []
    item_no: int | None = None
    name = ""
    for row in range(7, sheet.max_row + 1):
        number, title, length, pieces, grade = (sheet.cell(row, col).value for col in (2, 3, 6, 7, 8))
        if isinstance(number, int):
            item_no = number
        if title is not None and str(title).strip():
            name = str(title).strip()
        if length is None and pieces is None:
            continue
        if item_no is None or not name or length is None or pieces is None:
            raise ValueError(f"Заявка: неполная позиция в строке {row}")
        quantity = decimal_value(pieces)
        if quantity != int(quantity) or quantity <= 0:
            raise ValueError(f"Заявка: неверное количество в строке {row}")
        length_m = decimal_value(length)
        if length_m <= 0:
            raise ValueError(f"Заявка: неверная длина в строке {row}")
        lines.append(RequestLine(row, item_no, name, length_m, int(quantity), str(grade or "").upper()))
    if not lines:
        raise ValueError("В заявке не найдены строки материалов")
    return lines


def _invoice_number(text: str, fallback: str) -> str:
    match = re.search(r"СЧ[ЕЁ]Т(?:\s+НА\s+ОПЛАТУ)?\s*№\s*([А-ЯA-Z0-9-]+)", text, re.I)
    return match[1] if match else fallback


def _check_invoice_total(text: str, lines: list[InvoiceLine], filename: str) -> list[str]:
    """Сверить извлечённые строки с печатным итогом счёта, если он есть."""
    match = re.search(
        r"ВСЕГО\s+НАИМЕНОВАНИЙ\s+(\d+)\s*,\s*НА\s+СУММУ\s+([\d\s\u00a0.,]+)",
        text, re.I,
    )
    if not match:
        return [f"{filename}: итог счёта не найден; проверьте полноту извлечения"]
    expected_count = int(match[1])
    expected_amount = money(decimal_value(match[2]))
    actual_amount = money(sum((line.amount for line in lines), Decimal(0)))
    issues: list[str] = []
    if len(lines) != expected_count:
        issues.append(f"{filename}: строк извлечено {len(lines)} из {expected_count} по итогу счёта")
    if abs(actual_amount - expected_amount) > Decimal("1.00"):
        issues.append(f"{filename}: сумма извлечённых строк {actual_amount} ₽ отличается от итога "
                      f"счёта {expected_amount} ₽")
    return issues


def _table_kind(header: list[object]) -> str | None:
    labels = [str(cell or "").lower() for cell in header]
    if len(labels) >= 6 and "наименование товара" in labels[1] and "цена" in labels[-2]:
        return "evraz"
    if len(labels) >= 9 and "артикул" in labels[2] and "товары" in labels[3]:
        return "dip"
    if len(labels) >= 7 and "склад" in labels[4] and "товары" in labels[1]:
        return "corporation"
    return None


def _line_from_table(kind: str, row: list[object], supplier: str, invoice: str,
                     path: Path, page: int) -> InvoiceLine:
    number = int(str(row[0]).strip())
    if kind == "evraz":
        name, quantity, unit, price, amount = row[1:6]
        warehouse = ""
    elif kind == "dip":
        warehouse = str(row[1] or "").strip()
        name, quantity, unit, price, amount = (row[3], row[4], row[5], row[7], row[8])
    else:
        name, quantity, unit, warehouse, price, amount = row[1:7]
    name = re.sub(r"\s+", " ", str(name or "")).strip()
    quantity_d = decimal_value(quantity)
    price_d = decimal_value(price)
    amount_d = money(decimal_value(amount))
    unit_s = str(unit or "").strip().lower()
    length_m = length_from_name(name)
    notes: list[str] = []
    if unit_s in ("т", "тн", "тонна", "тонны"):
        mass_t = quantity_d
        expected = money(quantity_d * price_d)
    elif unit_s in ("шт", "штук"):
        kg_match = re.search(r"-\s*[кk]\s*(\d+(?:[.,]\d+)?)", name, re.I)
        mass_t = quantity_d * length_m * decimal_value(kg_match[1]) / 1000 if kg_match and length_m else None
        expected = money(quantity_d * price_d)
        if mass_t is None:
            notes.append("Масса неизвестна: нет проверяемых кг/м и длины")
    else:
        mass_t = None
        expected = amount_d
        notes.append(f"Неизвестная единица: {unit_s}")
    if abs(expected - amount_d) > Decimal("1.00"):
        notes.append(f"Цена × количество расходится с суммой на {abs(expected - amount_d)} ₽")
    price_per_t = money(amount_d / mass_t) if mass_t and mass_t > 0 else None
    return InvoiceLine(supplier, invoice, path, page, number, name, quantity_d, unit_s,
                       mass_t, amount_d, price_per_t, length_m, grade_tokens(name),
                       str(warehouse or "").strip(), tuple(notes))


def _windows_ocr(image_path: Path) -> list[dict[str, object]]:
    script = Path(__file__).with_name("windows_ocr.ps1")
    powershell = Path(os.environ.get("WINDIR", "C:/Windows")) / "System32/WindowsPowerShell/v1.0/powershell.exe"
    if os.name != "nt" or not powershell.is_file():
        raise RuntimeError("Для скана требуется Windows OCR")
    process = subprocess.run(
        [str(powershell), "-NoProfile", "-ExecutionPolicy", "Bypass", "-File", str(script), str(image_path)],
        capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=90,
    )
    if process.returncode != 0:
        raise RuntimeError("Windows OCR: " + process.stderr.strip()[:400])
    result = json.loads(process.stdout.strip())
    return result if isinstance(result, list) else [result]


def _read_scan_page(path: Path, page_number: int, supplier: str,
                    invoice_number: str) -> tuple[list[InvoiceLine], list[str]]:
    try:
        import pypdfium2 as pdfium
    except ImportError as exc:
        return [], ["Скан требует распознавания: установите pypdfium2"]
    with tempfile.TemporaryDirectory() as directory:
        document = pdfium.PdfDocument(str(path))
        page = document[page_number - 1]
        scale = 1800 / page.get_height()
        image = page.render(scale=scale).to_pil().convert("RGB")
        image_path = Path(directory) / "page.png"
        image.save(image_path)
        width, height = image.size
        try:
            blocks = _windows_ocr(image_path)
        except (RuntimeError, subprocess.TimeoutExpired, json.JSONDecodeError) as exc:
            return [], [f"Скан, стр. {page_number}: требуется распознавание или ручная проверка ({exc})"]
    text = " ".join(str(block.get("text", "")) for block in blocks[:12]).upper()
    if "МЕТАЛЛОТОРГ" not in text:
        return [], [f"Скан, стр. {page_number}: формат таблицы не распознан; требуется ручная проверка"]
    footer = min((float(b["y"]) for b in blocks if str(b.get("text", "")).strip().lower() == "итого"),
                 default=height * 0.73)
    top = height * 0.30
    bottom = min(footer, height * 0.73)

    def column(left: float, right: float, *, min_value: Decimal = Decimal("0")) -> list[tuple[float, Decimal]]:
        output: list[tuple[float, Decimal]] = []
        for block in blocks:
            x, y = float(block.get("x", -1)), float(block.get("y", -1))
            if not (left * width <= x < right * width and top < y < bottom):
                continue
            raw = str(block.get("text", ""))
            if not re.search(r"[0-9оОoO][.,][0-9оОoO]", raw):
                continue
            try:
                value = decimal_value(raw)
            except ValueError:
                continue
            if value >= min_value:
                output.append((y, value))
        return sorted(output)

    masses = column(0.39, 0.53)
    prices = column(0.53, 0.63, min_value=Decimal("1000"))
    amounts = column(0.80, 0.94, min_value=Decimal("1000"))
    if not masses or len(masses) != len(prices):
        return [], [f"Скан, стр. {page_number}: столбцы распознаны неполно "
                    f"(масса {len(masses)}, цена {len(prices)}, сумма {len(amounts)}); требуется ручная проверка"]
    lines: list[InvoiceLine] = []
    issues: list[str] = []
    for index, ((y_mass, mass), (y_price, price)) in enumerate(zip(masses, prices), 1):
        if abs(y_mass - y_price) > 15:
            issues.append(f"Скан, строка {index}: цена и масса находятся на разных строках")
            continue
        near = [(y, value) for y, value in amounts if abs(y - y_mass) <= 15]
        calculated = money(mass * price)
        if near:
            y_amount, amount = min(near, key=lambda part: abs(part[0] - y_mass))
            if abs(calculated - money(amount)) > Decimal("1.00"):
                issues.append(f"Скан, строка {index}: сумма OCR не сходится с массой × ценой; "
                              "использован расчёт, требуется проверка")
                amount = calculated
            else:
                amount = money(amount)
        else:
            amount = calculated
            issues.append(f"Скан, строка {index}: сумма не распознана; рассчитана по массе и цене, "
                          "требуется проверка")
        prev_y = masses[index - 2][0] if index > 1 else y_mass - 60
        next_y = masses[index][0] if index < len(masses) else y_mass + 60
        name_parts = [str(b.get("text", "")) for b in blocks
                      if float(b.get("x", width)) < width * 0.39
                      and (prev_y + y_mass) / 2 <= float(b.get("y", -1)) < (y_mass + next_y) / 2]
        raw_name = re.sub(r"^\s*\d{1,2}\s+", "", " ".join(name_parts)).strip()
        notes = ["OCR: проверить название, марку и длину по скану"]
        if not near or abs(calculated - money(near[0][1])) > Decimal("1.00"):
            notes.append("OCR: сумма рассчитана по массе и цене; проверить по скану")
        lines.append(InvoiceLine(
            supplier, invoice_number, path, page_number, index, raw_name,
            mass, "т", mass, money(amount), money(amount / mass) if mass > 0 else None,
            length_from_name(raw_name), grade_tokens(raw_name), "", tuple(notes),
        ))
    return lines, issues


def read_invoice(path: Path, *, supplier: str) -> Invoice:
    try:
        import pdfplumber
    except ImportError as exc:
        raise RuntimeError("Для сводной установите зависимости: pip install -r requirements.txt") from exc
    path = Path(path)
    if path.suffix.lower() != ".pdf":
        return Invoice(supplier, path.stem, path, (), (f"Формат {path.suffix} пока требует ручной проверки",))
    lines: list[InvoiceLine] = []
    issues: list[str] = []
    with pdfplumber.open(path) as document:
        pages_text = [page.extract_text() or "" for page in document.pages]
        invoice_number = _invoice_number("\n".join(pages_text), path.stem)
        for page_number, page in enumerate(document.pages, 1):
            found_table = False
            for table in page.extract_tables():
                if not table:
                    continue
                kind = _table_kind(table[0])
                if not kind:
                    continue
                found_table = True
                for row in table[1:]:
                    if not row or not str(row[0] or "").strip().isdigit():
                        continue
                    try:
                        lines.append(_line_from_table(kind, row, supplier, invoice_number, path, page_number))
                    except (ValueError, IndexError) as exc:
                        issues.append(f"{path.name}, стр. {page_number}, строка {row[0]}: {exc}")
            if not found_table and not pages_text[page_number - 1].strip():
                scanned, warnings = _read_scan_page(path, page_number, supplier, invoice_number)
                lines.extend(scanned)
                if not scanned and lines and page_number > 1 and all(
                    "формат таблицы не распознан" in warning for warning in warnings
                ):
                    # Заключительная страница с условиями оплаты, без товарных строк.
                    continue
                issues.extend(warnings)
            elif not found_table and pages_text[page_number - 1].strip():
                issues.append(f"{path.name}, стр. {page_number}: таблица счёта не распознана")
    if not lines:
        issues.append(f"{path.name}: ни одна позиция не извлечена; нужна ручная проверка")
    elif any(text.strip() for text in pages_text):
        issues.extend(_check_invoice_total("\n".join(pages_text), lines, path.name))
    return Invoice(supplier, invoice_number, path, tuple(lines), tuple(issues))


def _match_line(line: InvoiceLine, requests: list[RequestLine]) -> tuple[RequestLine | None, str]:
    key = line.key
    if key is None:
        return None, "Не сопоставлено: не удалось прочитать материал и размер"
    candidates = [request for request in requests if request.key == key]
    if not candidates:
        return None, "Не сопоставлено: материал или размер отсутствует в заявке"
    status: list[str] = []
    if line.length_m is not None:
        same_length = [request for request in candidates if request.length_m == line.length_m]
        if not same_length:
            return None, f"Длина отличается: {line.length_m} м"
        candidates = same_length
    elif len(candidates) > 1:
        status.append("Длина не указана в счёте")
    request = candidates[0]
    wanted_grades = grade_tokens(request.grade)
    wanted_grade = wanted_grades[0] if wanted_grades else None
    if not line.grades:
        status.append("Марка требует проверки")
    elif wanted_grade and wanted_grade not in line.grades:
        status.append("Марка отличается: " + "/".join("С" + grade for grade in line.grades))
    elif len(line.grades) > 1:
        status.append("Марка неоднозначна: " + "/".join("С" + grade for grade in line.grades))
    if line.mass_t is None:
        status.append("Масса неизвестна")
    if line.notes:
        status.extend(line.notes)
    return request, "; ".join(status) if status else "Совпадает"


def _digest(path: Path) -> str:
    checksum = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            checksum.update(chunk)
    return checksum.hexdigest()


def build_summary(request_path: Path, offer_sources: Iterable[OfferSource],
                  output_path: Path) -> SummaryResult:
    requests = read_request(Path(request_path))
    sources = list(offer_sources)
    if not sources:
        raise ValueError("Не выбраны счета поставщиков")
    suppliers = list(dict.fromkeys(source.supplier for source in sources))
    issues: list[str] = []
    allocations: list[Allocation] = []
    seen_files: set[tuple[str, str]] = set()
    seen_invoices: dict[tuple[str, str], str] = {}
    for source in sources:
        path = Path(source.path)
        if not path.is_file():
            issues.append(f"Не найден счёт: {path}")
            continue
        digest = _digest(path)
        if (source.supplier, digest) in seen_files:
            issues.append(f"Точная копия счёта пропущена: {path.name}")
            continue
        seen_files.add((source.supplier, digest))
        invoice = read_invoice(path, supplier=source.supplier)
        issues.extend(invoice.issues)
        invoice_key = (source.supplier, invoice.number)
        if invoice_key in seen_invoices and seen_invoices[invoice_key] != digest:
            issues.append(f"Повтор номера счёта {invoice.number} у {source.supplier}: {path.name}; "
                          "версия не включена до проверки")
            continue
        seen_invoices[invoice_key] = digest
        for line in invoice.lines:
            request, status = _match_line(line, requests)
            if request is None or status != "Совпадает":
                issues.append(f"{source.supplier}, {path.name}, строка {line.row_number}: {status}")
            allocations.append(Allocation(line, request, status))
    grouped: dict[tuple[int, str], list[Allocation]] = defaultdict(list)
    for allocation in allocations:
        if allocation.request is not None:
            grouped[(allocation.request.source_row, allocation.line.supplier)].append(allocation)
    offers: dict[tuple[int, str], CombinedOffer] = {}
    for (request_row, supplier), parts in grouped.items():
        request = next(request for request in requests if request.source_row == request_row)
        total_amount = sum((part.line.amount for part in parts), Decimal(0))
        masses = [part.line.mass_t for part in parts]
        total_mass = sum((mass for mass in masses if mass is not None), Decimal(0)) if all(
            mass is not None for mass in masses) else None
        average = money(total_amount / total_mass) if total_mass and total_mass > 0 else None
        statuses = list(dict.fromkeys(part.status for part in parts if part.status != "Совпадает"))
        offers[(request_row, supplier)] = CombinedOffer(
            supplier, request, tuple(part.line for part in parts), total_mass,
            money(total_amount), average, "; ".join(statuses) if statuses else "Совпадает",
        )
    result = SummaryResult(requests, suppliers, allocations, offers, issues, Path(output_path))
    _write_workbook(result)
    return result


def _write_workbook(result: SummaryResult) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    summary = workbook.active
    summary.title = "Сводная"
    detail = workbook.create_sheet("Партии")
    review = workbook.create_sheet("Проверка")
    detail.append(["Поставщик", "Счёт", "Файл", "Страница", "Строка", "Материал в счёте",
                   "Масса, т", "Сумма с НДС, ₽", "Цена за т, ₽", "Длина, м", "Марка", "Склад",
                   "Строка заявки", "Статус", "Примечание"])
    detail_rows: dict[InvoiceLine, int] = {}
    for allocation in result.allocations:
        line = allocation.line
        detail.append([line.supplier, line.invoice_number, line.source_file.name, line.page,
                       line.row_number, line.name, float(line.mass_t) if line.mass_t is not None else None,
                       float(line.amount), float(line.price_per_t) if line.price_per_t is not None else None,
                       float(line.length_m) if line.length_m is not None else None,
                       "/".join("С" + grade for grade in line.grades), line.warehouse,
                       allocation.request.source_row if allocation.request else None,
                       allocation.status, "; ".join(line.notes)])
        detail_rows[line] = detail.max_row
    review.append(["Проверить перед закупкой"])
    for issue in result.issues:
        review.append([issue])
    if not result.issues:
        review.append(["Замечаний нет"])

    last_col = 5 + 4 * len(result.suppliers)
    summary.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    summary.cell(1, 1, "Сводная предложений поставщиков")
    summary.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
    summary.cell(2, 1, "Цены за 1000 кг с НДС. Предложенные объёмы могут отличаться от заявки.")
    summary.merge_cells(start_row=3, start_column=1, end_row=3, end_column=last_col)
    summary.cell(3, 1, f"Строк заявки: {len(result.request_lines)}. Замечаний к проверке: {len(result.issues)}.")
    for col, title in enumerate(("№", "Материал", "Длина, м", "Кол-во, шт", "Марка"), 1):
        summary.cell(6, col, title)
    for index, supplier in enumerate(result.suppliers):
        start = 6 + index * 4
        summary.merge_cells(start_row=5, start_column=start, end_row=5, end_column=start + 3)
        summary.cell(5, start, supplier)
        for offset, title in enumerate(("Масса, т", "За 1 т, ₽", "Общая, ₽", "Проверка")):
            summary.cell(6, start + offset, title)
    for excel_row, request in enumerate(result.request_lines, 7):
        for col, value in enumerate((request.item_no, request.name, float(request.length_m),
                                     request.quantity_pcs, request.grade), 1):
            summary.cell(excel_row, col, value)
        for index, supplier in enumerate(result.suppliers):
            start = 6 + index * 4
            offer = result.offers.get((request.source_row, supplier))
            if not offer:
                summary.cell(excel_row, start + 3, "Нет предложения")
                continue
            rows = [detail_rows[line] for line in offer.source_lines]
            mass_refs = ",".join(f"'Партии'!G{row}" for row in rows)
            amount_refs = ",".join(f"'Партии'!H{row}" for row in rows)
            mass_cell = f"{get_column_letter(start)}{excel_row}"
            amount_cell = f"{get_column_letter(start + 2)}{excel_row}"
            summary.cell(excel_row, start, f"=SUM({mass_refs})" if offer.mass_t is not None else None)
            summary.cell(excel_row, start + 2, f"=SUM({amount_refs})")
            if offer.price_per_t is not None:
                summary.cell(excel_row, start + 1, f'=IF({mass_cell}>0,{amount_cell}/{mass_cell},"")')
            summary.cell(excel_row, start + 3, offer.status)
    total_row = 7 + len(result.request_lines)
    summary.cell(total_row, 2, "Итого по предложенным позициям")
    for index in range(len(result.suppliers)):
        start = 6 + index * 4
        mass_col = get_column_letter(start)
        sum_col = get_column_letter(start + 2)
        summary.cell(total_row, start, f"=SUM({mass_col}7:{mass_col}{total_row-1})")
        summary.cell(total_row, start + 2, f"=SUM({sum_col}7:{sum_col}{total_row-1})")
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    summary.freeze_panes = "F7"
    detail.freeze_panes = "G2"
    summary.row_dimensions[1].height = 34
    summary.row_dimensions[2].height = 24
    summary.row_dimensions[3].height = 24
    summary.row_dimensions[5].height = 30
    summary.row_dimensions[6].height = 32
    summary.auto_filter.ref = f"A6:E{total_row-1}"
    detail.auto_filter.ref = f"A1:O{detail.max_row}"
    summary.column_dimensions["B"].width = 30
    for col in ("A", "C", "D", "E"):
        summary.column_dimensions[col].width = 13
    for index in range(len(result.suppliers)):
        start = 6 + index * 4
        for offset, width in enumerate((14, 17, 18, 38)):
            summary.column_dimensions[get_column_letter(start + offset)].width = width
    for col, width in {"A": 28, "B": 18, "C": 38, "D": 12, "E": 10, "F": 55,
                       "G": 14, "H": 18, "I": 17, "J": 12, "K": 16, "L": 20,
                       "M": 14, "N": 45, "O": 50}.items():
        detail.column_dimensions[col].width = width
    review.column_dimensions["A"].width = 115
    navy = PatternFill("solid", fgColor="17324D")
    blue = PatternFill("solid", fgColor="DCEAF7")
    warning = PatternFill("solid", fgColor="FFF2CC")
    for sheet, header_rows in ((summary, (1, 5, 6)), (detail, (1,)), (review, (1,))):
        for row in header_rows:
            for cell in sheet[row]:
                if cell.__class__.__name__ == "MergedCell":
                    continue
                cell.fill = navy if row == 1 else blue
                cell.font = Font(color="FFFFFF" if row == 1 else "17324D", bold=True)
                cell.alignment = Alignment(vertical="center", wrap_text=True)
    for row in summary.iter_rows(min_row=7, max_row=total_row):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if cell.column in (1, 4):
                cell.number_format = "#,##0"
            elif cell.column == 3:
                cell.number_format = "0.##"
        for index in range(len(result.suppliers)):
            mass_cell = row[5 + index * 4]
            price_cell = row[6 + index * 4]
            amount_cell = row[7 + index * 4]
            mass_cell.number_format = "#,##0.000"
            price_cell.number_format = "#,##0.00"
            amount_cell.number_format = "#,##0.00"
            status_cell = row[5 + index * 4 + 3]
            if status_cell.value not in ("Совпадает", "Нет предложения", None):
                status_cell.fill = warning
        summary.row_dimensions[row[0].row].height = 38 if any(
            len(str(row[5 + index * 4 + 3].value or "")) > 55
            for index in range(len(result.suppliers))
        ) else 27
    for row in detail.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        row[6].number_format = "#,##0.000"
        for col in (8, 9):
            row[col - 1].number_format = "#,##0.00"
    for row in review.iter_rows(min_row=2):
        row[0].alignment = Alignment(wrap_text=True, vertical="top")
    result.output_path.parent.mkdir(parents=True, exist_ok=True)
    temporary = result.output_path.with_name(result.output_path.stem + ".tmp.xlsx")
    try:
        workbook.save(temporary)
        temporary.replace(result.output_path)
    finally:
        temporary.unlink(missing_ok=True)
