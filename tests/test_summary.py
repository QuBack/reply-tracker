from __future__ import annotations

import os
import tempfile
import unittest
from decimal import Decimal
from pathlib import Path

from automation.summary import OfferSource, _check_invoice_total, build_summary, read_invoice, read_request


FILES = {
    "request": "Заявка 03.09.2026г Липяги Раб. башня (Прокат) (1) — копия.xlsx",
    "evraz_317": "Счет на оплату Счет - Оферта № 218317 от 03.09.2026.pdf",
    "evraz_349": "Счет на оплату Счет - Оферта № 218349 от 03.09.2026.pdf",
    "evraz_331": "Счет на оплату Счет - Оферта № 218331 от 03.09.2026.pdf",
    "dip": "ДИП Счет.pdf",
    "corporation": "40675.pdf",
    "scan": "Металлоторг_ВсеЛисты (14).pdf",
}


def fixtures() -> dict[str, Path]:
    root = Path(os.environ.get("SUMMARY_FIXTURES_DIR", Path.home() / "Downloads"))
    return {key: root / filename for key, filename in FILES.items()}


class SummaryFixtureTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.files = fixtures()
        missing = [path.name for path in cls.files.values() if not path.is_file()]
        if missing:
            raise unittest.SkipTest("Нет тестовых документов: " + ", ".join(missing))

    def test_request_keeps_all_29_length_rows(self) -> None:
        lines = read_request(self.files["request"])
        self.assertEqual(len(lines), 29)
        self.assertEqual(len({line.item_no for line in lines}), 19)
        self.assertEqual([(line.length_m, line.quantity_pcs) for line in lines if line.item_no == 18],
                         [(Decimal("12"), 26), (Decimal("8"), 1)])
        self.assertTrue(all(line.name and line.grade for line in lines))

    def test_text_invoice_extracts_source_and_per_piece_mass(self) -> None:
        invoice = read_invoice(self.files["corporation"], supplier="Корпорация МеталлИнвест")
        self.assertEqual(invoice.number, "40675")
        self.assertEqual(len(invoice.lines), 17)
        line = next(line for line in invoice.lines if "120х120х4" in line.name)
        self.assertEqual(line.mass_t, Decimal("1.368"))
        self.assertEqual(line.amount, Decimal("94392.00"))
        self.assertEqual(line.price_per_t, Decimal("69000"))
        self.assertEqual(line.warehouse, "ф-л_Оскол")
        self.assertEqual(line.page, 1)

    def test_text_invoice_total_catches_missing_rows(self) -> None:
        for key in ("evraz_317", "evraz_331", "evraz_349", "dip", "corporation"):
            invoice = read_invoice(self.files[key], supplier="Поставщик")
            self.assertFalse(invoice.issues, f"{key}: {invoice.issues}")
        invoice = read_invoice(self.files["corporation"], supplier="Поставщик")
        issues = _check_invoice_total(
            "Всего наименований 17, на сумму 5 127 886,35 RUB",
            list(invoice.lines[:-1]), invoice.source_file.name,
        )
        self.assertTrue(any("строк извлечено" in issue for issue in issues))
        self.assertTrue(any("сумма извлечённых" in issue for issue in issues))

    def test_weighted_price_across_evraz_invoices(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            result = build_summary(
                self.files["request"],
                [
                    OfferSource(self.files["evraz_317"], "ЕВРАЗ Маркет"),
                    OfferSource(self.files["evraz_349"], "ЕВРАЗ Маркет"),
                ],
                Path(directory) / "summary.xlsx",
            )
            offer = result.find_offer(item_no=19, length_m=Decimal("12"), supplier="ЕВРАЗ Маркет")
            self.assertIsNotNone(offer)
            self.assertEqual(offer.mass_t, Decimal("3.456"))
            self.assertEqual(offer.amount, Decimal("304281.60"))
            self.assertEqual(offer.price_per_t, Decimal("88044.44"))
            self.assertEqual(len(offer.source_lines), 2)

    def test_grade_and_length_mismatch_are_visible(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            result = build_summary(
                self.files["request"],
                [OfferSource(self.files["evraz_331"], "ЕВРАЗ Маркет")],
                Path(directory) / "summary.xlsx",
            )
            offer = result.find_offer(item_no=18, length_m=Decimal("12"), supplier="ЕВРАЗ Маркет")
            self.assertIsNotNone(offer)
            self.assertIn("марка", offer.status.lower())
            self.assertIsNone(result.find_offer(item_no=18, length_m=Decimal("8"), supplier="ЕВРАЗ Маркет"))

    def test_repeated_pdf_does_not_double_count(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            source = OfferSource(self.files["evraz_317"], "ЕВРАЗ Маркет")
            result = build_summary(
                self.files["request"], [source, source], Path(directory) / "summary.xlsx"
            )
            offer = result.find_offer(item_no=19, length_m=Decimal("12"), supplier="ЕВРАЗ Маркет")
            self.assertEqual(offer.amount, Decimal("133785.60"))
            self.assertEqual(len(offer.source_lines), 1)

    def test_xlsx_has_29_rows_formulas_and_provenance(self) -> None:
        from openpyxl import load_workbook

        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "summary.xlsx"
            build_summary(
                self.files["request"],
                [OfferSource(self.files["evraz_317"], "ЕВРАЗ Маркет")],
                output,
            )
            workbook = load_workbook(output, data_only=False)
            self.assertIn("Сводная", workbook.sheetnames)
            self.assertIn("Партии", workbook.sheetnames)
            summary = workbook["Сводная"]
            self.assertEqual(sum(isinstance(summary.cell(row, 1).value, int)
                                 for row in range(1, summary.max_row + 1)), 29)
            self.assertTrue(any(cell.data_type == "f" for row in summary for cell in row))
            detail = workbook["Партии"]
            self.assertTrue(any("218317" in str(cell.value or "") for row in detail for cell in row))

    def test_service_uses_only_selected_campaign_files(self) -> None:
        from openpyxl import load_workbook

        from automation.database import Database
        from automation.paths import AppPaths
        from automation.service import AppService

        class NoCredentials:
            pass

        with tempfile.TemporaryDirectory() as directory:
            paths = AppPaths.from_root(Path(directory))
            database = Database(paths.database)
            service = AppService(paths, database, NoCredentials())
            supplier_id = database.save_supplier("ЕВРАЗ Маркет", "evraz@example.ru")
            campaign_id = database.create_campaign("Тест", "Запрос", "Текст", None, [supplier_id])
            recipient_id = int(database.list_recipients(campaign_id)[0]["id"])
            with database.connect() as connection:
                connection.execute("UPDATE campaigns SET status = 'ready' WHERE id = ?", (campaign_id,))
            incoming_id, created = database.insert_incoming(
                mailbox="buyer@mail.ru", folder="INBOX", uid_validity="1", imap_uid=1,
                campaign_id=campaign_id, recipient_id=recipient_id,
                message_id="<reply@example.ru>", in_reply_to="", sender_email="evraz@example.ru",
                subject="Ответ", received_at=None, body_text="", raw_eml_path="",
                match_method="manual", needs_review=False,
            )
            self.assertTrue(created)
            pdf = self.files["evraz_317"]
            database.add_incoming_attachment(
                incoming_message_id=incoming_id, filename=pdf.name,
                content_type="application/pdf", size=pdf.stat().st_size,
                sha256="sample-pdf", path=str(pdf), is_allowed=True, duplicate_of_id=None,
            )
            result = service.create_summary(campaign_id, self.files["request"])
            output = Path(str(result.details["path"]))
            self.assertTrue(output.is_file())
            saved_request = Path(database.get_campaign(campaign_id)["request_path"])
            self.assertTrue(saved_request.is_file())
            self.assertNotEqual(saved_request, self.files["request"])
            repeated = service.create_summary(campaign_id)
            self.assertTrue(Path(str(repeated.details["path"])).is_file())
            workbook = load_workbook(output, read_only=True)
            try:
                self.assertEqual(workbook["Сводная"]["F5"].value, "ЕВРАЗ Маркет")
            finally:
                workbook.close()

    @unittest.skipUnless(os.name == "nt", "Windows OCR проверяется на Windows")
    def test_scan_recovers_weighted_price_with_review_flags(self) -> None:
        invoice = read_invoice(self.files["scan"], supplier="Металлоторг")
        self.assertEqual(len(invoice.lines), 19)
        self.assertEqual(sum((line.amount for line in invoice.lines), Decimal(0)),
                         Decimal("5705387.20"))
        with tempfile.TemporaryDirectory() as directory:
            result = build_summary(
                self.files["request"],
                [OfferSource(self.files["scan"], "Металлоторг")],
                Path(directory) / "summary.xlsx",
            )
            offer = result.find_offer(item_no=18, length_m=Decimal("12"), supplier="Металлоторг")
            self.assertEqual(offer.mass_t, Decimal("21.725"))
            self.assertEqual(offer.price_per_t, Decimal("99540.00"))
            self.assertIn("провер", offer.status.lower())


if __name__ == "__main__":
    unittest.main()
